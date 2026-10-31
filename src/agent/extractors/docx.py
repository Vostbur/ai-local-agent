"""Экстрактор текста из файлов DOCX."""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from agent.extractors.base import ExtractedDocument, Segment, TextExtractor

class DocxExtractor(TextExtractor):
    SUFFIXES = {".docx"}

    def extract(self, path: Path) -> ExtractedDocument | None:
        try:
            doc = DocxDocument(path)
        except Exception:
            return None

        segments = []
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                # Используем стиль параграфа как секцию, если он есть
                section = para.style.name if para.style else ""
                segments.append(Segment(text=text, section=section))

        return ExtractedDocument(
            path=path,
            segments=segments,
            metadata={"format": "docx", "paragraphs": len(doc.paragraphs)},
        )